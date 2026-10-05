"""Show typed multi-sheet queries and cache reuse after an unrelated sheet edit."""

import argparse
from pathlib import Path

from openpyxl import Workbook as Excel

from sheetjet import Workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("benchmark-output/multisheet-demo"))
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    source = args.output / "periods.xlsx"
    revised = args.output / "revised.xlsx"
    cache = args.output / "cache"
    excel = Excel()
    january = excel.active
    january.title = "January"
    for row in [["Region", "Revenue"], ["APAC", 125.5], ["EMEA", 80]]:
        january.append(row)
    february = excel.create_sheet("February")
    for row in [["Revenue", "Region"], [150.25, "APAC"], [95.75, "EMEA"]]:
        february.append(row)
    excel.create_sheet("Assumptions").append(["Growth", 0.08])
    excel.save(source)
    ranges = {"January": "A1:B3", "February": "A1:B3"}
    schema = {"Region": "VARCHAR", "Revenue": "DECIMAL(18,2)"}
    sql = 'SELECT "Region", sum("Revenue") FROM Sales GROUP BY 1 ORDER BY 1'
    with Workbook(source, cache) as book:
        book.query_engine.load_sheets("Sales", ranges, schema)
        answer = book.query_engine.query(sql)
        assert answer["rows"] == [["APAC", "275.75"], ["EMEA", "175.75"]]
        print("Reordered headers align by name:", book.serialize(answer))
        # Opt into the temporary table only when planning repeated scans.
        book.query_engine.materialize("Sales")
        assert book.query_engine.query(sql) == answer
        book.patch_cells(
            [
                {
                    "operation": "SET_VALUE",
                    "sheet": "Assumptions",
                    "cell": "B1",
                    "expected": 0.08,
                    "value": 0.12,
                }
            ],
            revised,
            overwrite=True,
        )
    with Workbook(revised, cache) as book:
        book.query_engine.load_sheets("Sales", ranges, schema)
        assert book.metrics.counts["persistent_table_hits"] == 2
        assert book.query_engine.query(sql) == answer
        print("Both period projections reused after editing Assumptions:", book.metrics.snapshot())


if __name__ == "__main__":
    main()
