import pytest
from openpyxl import Workbook as Excel
from openpyxl.utils import get_column_letter

from sheetjet import SheetJetError, Workbook
from sheetjet.core import Budget


def test_wide_headers_stay_local_with_tiny_result_budget(tmp_path):
    source = tmp_path / "wide.xlsx"
    excel = Excel()
    excel.active.append([f"Field{i}" for i in range(2500)])
    excel.active.append(list(range(2500)))
    excel.save(source)
    with Workbook(
        source, tmp_path / "cache", Budget(max_cells=2, max_rows=1, max_chars=100)
    ) as book:
        book.query_engine.load(
            "Wide",
            sheet="Sheet",
            ref=f"A1:{get_column_letter(2500)}2",
            columns=["Field0", "Field2499"],
            schema={"Field0": "BIGINT", "Field2499": "BIGINT"},
        )
        assert book.metrics.counts.get("characters_exposed", 0) == 0
        assert book.query_engine.query("SELECT Field0,Field2499 FROM Wide")["rows"] == [[0, 2499]]


def test_numeric_headers_remain_invalid(tmp_path):
    source = tmp_path / "numeric.xlsx"
    excel = Excel()
    excel.active.append([123, "Amount"])
    excel.active.append([1, 2])
    excel.save(source)
    with (
        Workbook(source, tmp_path / "cache") as book,
        pytest.raises(SheetJetError, match="Headers"),
    ):
        book.query_engine.load("Bad", sheet="Sheet", ref="A1:B2")
