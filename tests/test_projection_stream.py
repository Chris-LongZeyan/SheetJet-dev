import random
from zipfile import ZipFile

import pytest
from lxml import etree as ET
from openpyxl import Workbook as Excel

from sheetjet import SheetJetError, UnsupportedOperation, Workbook
from sheetjet.core import NS, Q, address


def workbook_xml(tmp_path, raw):
    base = tmp_path / "base.xlsx"
    Excel().save(base)
    target = tmp_path / "source.xlsx"
    with ZipFile(base) as source, ZipFile(target, "w") as output:
        for info in source.infolist():
            output.writestr(
                info, raw if info.filename == "xl/worksheets/sheet1.xml" else source.read(info)
            )
    return target


@pytest.mark.parametrize("seed", [0, 17, 92])
@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_stream_matches_tree_decoder_on_sparse_mixed_cells(tmp_path, cache, seed, encoding):
    rng = random.Random(seed)
    root = ET.Element(Q + "worksheet", nsmap={"m": NS})
    data = ET.SubElement(root, Q + "sheetData")
    for rn in sorted(rng.sample(range(1, 151), 90)):
        row = ET.SubElement(data, Q + "row", r=str(rn))
        for col in sorted(rng.sample(range(1, 7), rng.randrange(7))):
            node = ET.SubElement(row, Q + "c", r=address(rn, col))
            kind = rng.choice(["n", "b", "inlineStr", "str", "blank", "formula"])
            if kind in {"n", "formula"}:
                if kind == "formula":
                    ET.SubElement(node, Q + "f").text = "1/3"
                ET.SubElement(node, Q + "v").text = "0.10000000000000000001"
            elif kind == "b":
                node.set("t", "b")
                ET.SubElement(node, Q + "v").text = str(rng.randrange(2))
            elif kind == "str":
                node.set("t", "str")
                ET.SubElement(node, Q + "v").text = "<>& 日本語\n spaced "
            elif kind == "inlineStr":
                node.set("t", "inlineStr")
                inline = ET.SubElement(node, Q + "is")
                for text in ["a & b", " 中 ", "line\nbreak"]:
                    run = ET.SubElement(inline, Q + "r")
                    ET.SubElement(run, Q + "t").text = ET.CDATA(text)
                phonetic = ET.SubElement(inline, Q + "rPh", sb="0", eb="1")
                ET.SubElement(phonetic, Q + "t").text = "exclude phonetics"
    source = workbook_xml(tmp_path, ET.tostring(root, encoding=encoding, xml_declaration=True))
    with Workbook(source, cache) as book:
        expected = []
        for row in root.find(Q + "sheetData"):
            rn = int(row.get("r"))
            if 2 <= rn <= 145:
                cells = {c.get("r"): c for c in row}
                expected.append(
                    (
                        rn,
                        [
                            book.package.cell(
                                cells[address(rn, c)],
                                values_only=True,
                                allow_cached_formulas=True,
                                numeric_text=True,
                            )
                            if address(rn, c) in cells
                            else None
                            for c in [4, 1, 3]
                        ],
                    )
                )
        assert list(book.package.projected_rows("Sheet", "A2:F145", [4, 1, 3], True)) == expected


def test_stream_ignores_foreign_namespace_cells_and_unselected_errors(tmp_path, cache):
    raw = f'''<worksheet xmlns="{NS}" xmlns:x="urn:unrelated"><sheetData>
      <row r="2"><x:c r="A2"><x:v>evil</x:v></x:c>
      <c r="A2" t="inlineStr"><is><t>real</t></is></c>
      <c r="B2" t="e"><v>#DIV/0!</v></c></row>
      <row r="3"><c r="A3"><f>1+1</f><v>2</v></c></row>
    </sheetData></worksheet>'''.encode()
    with Workbook(workbook_xml(tmp_path, raw), cache) as book:
        assert list(book.package.projected_rows("Sheet", "A2:A2", [1])) == [(2, ["real"])]
        with pytest.raises(SheetJetError, match="Excel error"):
            list(book.package.projected_rows("Sheet", "B2:B2", [2]))
        with pytest.raises(SheetJetError, match="Formula inputs"):
            list(book.package.projected_rows("Sheet", "A3:A3", [1]))
        assert list(book.package.projected_rows("Sheet", "A3:A3", [1], True)) == [(3, ["2"])]


@pytest.mark.parametrize(
    "content,match",
    [
        ('<row r="2"><c r="A3"><v>1</v></c></row>', "does not match"),
        ('<row r="2"/><row r="2"/>', "increasing"),
        ('<row r="2"><c r="A2"/><c r="A2"/></row>', "Duplicate projected cell"),
    ],
)
def test_stream_rejects_ambiguous_coordinates(tmp_path, cache, content, match):
    raw = f'<worksheet xmlns="{NS}"><sheetData>{content}</sheetData></worksheet>'.encode()
    with (
        Workbook(workbook_xml(tmp_path, raw), cache) as book,
        pytest.raises(SheetJetError, match=match),
    ):
        list(book.package.projected_rows("Sheet", "A1:A5", [1]))


def test_stream_rejects_dtd(tmp_path, cache):
    raw = f'<!DOCTYPE worksheet [<!ENTITY x "expanded">]><worksheet xmlns="{NS}"><sheetData/></worksheet>'.encode()
    with (
        Workbook(workbook_xml(tmp_path, raw), cache) as book,
        pytest.raises(UnsupportedOperation, match="DTD"),
    ):
        list(book.package.projected_rows("Sheet", "A1:A5", [1]))
