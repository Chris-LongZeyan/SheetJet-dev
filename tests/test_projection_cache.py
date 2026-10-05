import pytest
from openpyxl import Workbook as Excel

from sheetjet import SheetJetError, Workbook


def test_persistent_projection_and_sheet_invalidation(model, cache, tmp_path):
    args = {"columns": ["Region", "Revenue"], "schema": {"Region": "VARCHAR", "Revenue": "BIGINT"}}
    with Workbook(model, cache) as w:
        w.query_engine.load("Sales", **args)
        expected = w.query_engine.query("SELECT sum(Revenue) FROM Sales")["rows"]
    with Workbook(model, cache) as w:
        assert w.query_engine.load("Sales", **args)["persistent_cache"]
        assert w.metrics.counts["cells_scanned"] <= 5  # header only, no data rows
        assert w.query_engine.query("SELECT sum(Revenue) FROM Sales")["rows"] == expected
        other = tmp_path / "other.xlsx"
        w.patch_cells(
            [{"operation": "SET_VALUE", "sheet": "Assumptions", "cell": "B1", "value": 99}], other
        )
    with Workbook(other, cache) as w:
        assert w.query_engine.load("Sales", **args)["persistent_cache"]
        changed = tmp_path / "changed.xlsx"
        w.patch_cells(
            [{"operation": "SET_VALUE", "sheet": "Transactions", "cell": "C2", "value": 999}],
            changed,
        )
    with Workbook(changed, cache) as w:
        assert not w.query_engine.load("Sales", **args)["persistent_cache"]
        assert w.query_engine.query("SELECT sum(Revenue) FROM Sales")["rows"] == [
            [expected[0][0] + 979]
        ]


def test_corrupt_cache_recovers_and_opt_out(model, cache):
    args = {"columns": ["Revenue"], "schema": {"Revenue": "BIGINT"}}
    with Workbook(model, cache) as w:
        w.query_engine.load("Sales", **args)
    files = list(cache.glob("projection-*.parquet"))
    assert len(files) == 1
    files[0].write_bytes(b"invalid parquet")
    with Workbook(model, cache) as w:
        assert not w.query_engine.load("Sales", **args)["persistent_cache"]
        assert w.query_engine.query("SELECT sum(Revenue) FROM Sales")["rows"] == [[51500]]
    with Workbook(model, cache) as w:
        assert not w.query_engine.load("Sales", persistent_cache=False, **args)["persistent_cache"]


def test_multi_sheet_union_reordered_headers_and_provenance(tmp_path, cache):
    path = tmp_path / "multi.xlsx"
    b = Excel()
    a = b.active
    a.title = "January"
    for row in [["Region", "Revenue"], ["APAC", 10], ["EMEA", 20]]:
        a.append(row)
    a = b.create_sheet("February's data")
    for row in [["Revenue", "Region"], [30, "APAC"], [40, "EMEA"]]:
        a.append(row)
    a = b.create_sheet("Irrelevant")
    a["A1"] = "Never stage this sheet"
    b.save(path)
    ranges = {"January": "A1:B3", "February's data": "A1:B3"}
    schema = {"Region": "VARCHAR", "Revenue": "BIGINT"}
    with Workbook(path, cache) as w:
        w.query_engine.load_sheets("AllSales", ranges, schema)
        assert w.query_engine.query(
            "SELECT Region,sum(Revenue) FROM AllSales GROUP BY 1 ORDER BY 1"
        )["rows"] == [["APAC", 40], ["EMEA", 60]]
        assert w.query_engine.query("SELECT DISTINCT _sheet FROM AllSales ORDER BY 1")["rows"] == [
            ["February's data"],
            ["January"],
        ]
        with pytest.raises((SheetJetError, KeyError)):
            w.query_engine.load_sheets("AllSales", {"January": "A1:B3", "Missing": "A1:B3"}, schema)
        assert w.query_engine.query("SELECT sum(Revenue) FROM AllSales")["rows"] == [[100]]
    with Workbook(path, cache) as w:
        w.query_engine.load_sheets("AllSales", ranges, schema)
        assert w.metrics.counts["persistent_table_hits"] == 2


def test_cache_keeps_formula_policy_and_exact_decimals(tmp_path, cache):
    from zipfile import ZipFile

    from lxml import etree as ET

    from sheetjet.core import Q

    path = tmp_path / "formula.xlsx"
    b = Excel()
    b.active.append(["Amount"])
    b.active.append(["=0.1+0.2"])
    b.save(path)
    temp = tmp_path / "cached.xlsx"
    with ZipFile(path) as a, ZipFile(temp, "w") as out:
        for info in a.infolist():
            raw = a.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(raw)
                root.find(f'.//{Q}c[@r="A2"]/{Q}v').text = "0.10000000000000000001"
                raw = ET.tostring(root)
            out.writestr(info, raw)
    args = {"sheet": "Sheet", "ref": "A1:A2", "schema": {"Amount": "DECIMAL(38,20)"}}
    for _ in range(2):
        with Workbook(temp, cache) as w:
            w.query_engine.load("Amounts", allow_cached_formulas=True, **args)
            assert w.query_engine.query("SELECT Amount FROM Amounts")["rows"] == [
                ["0.10000000000000000001"]
            ]
            with pytest.raises(SheetJetError, match="Formula inputs"):
                w.query_engine.load("Unsafe", **args)
