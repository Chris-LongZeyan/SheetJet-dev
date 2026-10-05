import hashlib
import json
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from sheetjet import SheetJetError, Workbook


def edit(sheet, cell, value, **extra):
    return {"operation": "SET_VALUE", "sheet": sheet, "cell": cell, "value": value, **extra}


def test_surgical_patch_preserves_opaque_parts(model, cache, tmp_path):
    original_hash = hashlib.sha256(model.read_bytes()).hexdigest()
    output = tmp_path / "edited.xlsx"
    with Workbook(model, cache) as w:
        report = w.patch_cells([edit("Assumptions", "B1", 0.12, expected=0.08)], output)
    assert report["requires_recalculation"]
    assert hashlib.sha256(model.read_bytes()).hexdigest() == original_hash
    with ZipFile(model) as a, ZipFile(output) as b:
        for name in a.namelist():
            if name not in report["changed_parts"]:
                assert a.read(name) == b.read(name)
        before = a.read("xl/worksheets/sheet2.xml")
        after = b.read("xl/worksheets/sheet2.xml")
        # Excluding the edited cell, every XML byte survives unchanged.
        import re

        def remove(x):
            return re.sub(rb'<c\b[^>]*\br="B1"[^>]*>.*?</c>', b"", x)

        assert remove(before) == remove(after)
    b = load_workbook(output)
    assert b["Assumptions"]["B1"].value == 0.12
    assert b["Assumptions"]["B1"].number_format == "0.0%"
    assert b["Transactions"]["E2"].value == "=C2-D2"
    assert len(b["Transactions"]._charts) == 1
    assert (
        json.loads(output.with_suffix(".xlsx.sheetjet.json").read_text())["changes"][0]["before"][
            "value"
        ]
        == 0.08
    )
    b.close()


def test_insert_cells_rows_and_copy_formula(model, cache, tmp_path):
    output = tmp_path / "expanded.xlsx"
    ops = [
        edit("Assumptions", "C1", "=literal string"),
        edit("Assumptions", "B3", False),
        edit("Assumptions", "A30", 17),
        {"operation": "COPY_FORMULA", "sheet": "Assumptions", "source_cell": "B8", "cell": "C8"},
        {"operation": "COPY_STYLE", "sheet": "Assumptions", "source_cell": "B1", "cell": "C1"},
    ]
    with Workbook(model, cache) as w:
        w.patch_cells(ops, output)
    b = load_workbook(output)
    a = b["Assumptions"]
    assert a["C1"].value == "=literal string" and a["C1"].data_type == "s"
    assert a["C1"].number_format == "0.0%"
    assert a["B3"].value is False
    assert a["A30"].value == 17
    assert a["C8"].value == "=C1*100"
    assert a.max_row == 30
    b.close()


@pytest.mark.parametrize(
    "ops,match",
    [
        ([edit("Assumptions", "B1", 1, expected=9)], "Precondition"),
        ([edit("Assumptions", "B4", 1)], "merged"),
        ([edit("Transactions", "A1", "Changed")], "header"),
        ([edit("Assumptions", "B1", 1), edit("Assumptions", "B1", 2)], "Multiple"),
        ([{"operation": "INSERT_ROWS", "sheet": "Assumptions", "cell": "A1"}], "preservation-safe"),
        ([edit("Assumptions", "B1", float("nan"))], "NaN"),
        (
            [{"operation": "SET_STYLE", "sheet": "Assumptions", "cell": "A1", "style_id": 99999}],
            "Style",
        ),
    ],
)
def test_reject_unsafe_without_output(model, cache, tmp_path, ops, match):
    out = tmp_path / "bad.xlsx"
    with Workbook(model, cache) as w, pytest.raises(SheetJetError, match=match):
        w.patch_cells(ops, out)
    assert not out.exists()


def test_write_range_empty_sheet(tmp_path, cache):
    from openpyxl import Workbook as Excel

    path = tmp_path / "empty.xlsx"
    Excel().save(path)
    out = tmp_path / "written.xlsx"
    with Workbook(path, cache) as w:
        w.write_range("Sheet", "B2", [[1, 2], [3, 4]], out)
    b = load_workbook(out)
    assert b.active["B2"].value == 1
    assert b.active["C3"].value == 4
    b.close()


def test_source_and_existing_output_protected(model, cache, tmp_path):
    with Workbook(model, cache) as w:
        with pytest.raises(SheetJetError, match="distinct"):
            w.patch_cells([edit("Assumptions", "B1", 1)], model)
        out = tmp_path / "exists.xlsx"
        out.write_bytes(b"do not touch")
        with pytest.raises(SheetJetError, match="exists"):
            w.patch_cells([edit("Assumptions", "B1", 1)], out)
        assert out.read_bytes() == b"do not touch"


def test_patch_preserves_escaped_namespace_uri(model, cache, tmp_path):
    source = tmp_path / "escaped.xlsx"
    with ZipFile(model) as original, ZipFile(source, "w") as rewritten:
        for info in original.infolist():
            data = original.read(info)
            if info.filename == "xl/worksheets/sheet2.xml":
                data = data.replace(
                    b"<worksheet ",
                    b'<worksheet xmlns:extra="urn:sheetjet?a=1&amp;b=2" ',
                    1,
                )
            rewritten.writestr(info, data)
    output = tmp_path / "escaped-patched.xlsx"
    with Workbook(source, cache) as book:
        book.patch_cells([edit("Assumptions", "B1", 0.15)], output)
    loaded = load_workbook(output)
    assert loaded["Assumptions"]["B1"].value == 0.15
    loaded.close()
    with ZipFile(output) as package:
        assert b'xmlns:extra="urn:sheetjet?a=1&amp;b=2"' in package.read("xl/worksheets/sheet2.xml")
