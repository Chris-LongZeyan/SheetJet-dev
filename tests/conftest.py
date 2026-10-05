import pytest
from openpyxl import Workbook as ExcelWorkbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table


@pytest.fixture
def model(tmp_path):
    path = tmp_path / "model.xlsx"
    book = ExcelWorkbook()
    s = book.active
    s.title = "Transactions"
    s.append(["Region", "Customer", "Revenue", "Cost", "Profit"])
    for i in range(2, 102):
        s.append(["APAC" if i % 2 == 0 else "EMEA", f"C{i % 5}", i * 10, i * 6, f"=C{i}-D{i}"])
    s.add_table(Table(displayName="Sales", ref="A1:E101"))
    s.freeze_panes = "C2"
    s["C2"].number_format = '"$"#,##0.00'
    s["C2"].fill = PatternFill("solid", fgColor="FF123456")
    s["C2"].font = Font(name="Arial", bold=True, color="FFFFFFFF")
    s.row_dimensions[10].hidden = True
    s.column_dimensions["D"].hidden = True
    s.row_dimensions[2].height = 23
    s.column_dimensions["A"].width = 18
    a = book.create_sheet("Assumptions")
    a.append(["China revenue growth assumption", 0.08])
    a.append(["EBITDA margin", 0.25])
    a.merge_cells("A4:C4")
    a["A4"] = "Merged title"
    a["F20"] = "Unindexed interior phrase"
    a["G21"] = "Need full search"
    a["B1"].number_format = "0.0%"
    a["B8"] = "=B1*100"
    hidden = book.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Secret local-only label"
    book.defined_names.add(DefinedName("Growth", attr_text="'Assumptions'!$B$1"))
    chart = BarChart()
    chart.add_data(Reference(s, min_col=3, max_col=3, min_row=1, max_row=6), titles_from_data=True)
    s.add_chart(chart, "H2")
    book.save(path)
    return path


@pytest.fixture
def cache(tmp_path):
    return tmp_path / "cache"
