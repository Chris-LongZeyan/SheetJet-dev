import json

from openpyxl import Workbook as Excel
from openpyxl import load_workbook

from benchmarks.skill_trials import cell_fingerprints, expected_changes, judge
from sheetjet import Workbook


def test_judge_checks_saved_values_styles_and_unrequested_cells(tmp_path):
    source, output, report = (tmp_path / name for name in ["in.xlsx", "out.xlsx", "result.json"])
    excel = Excel()
    sheet = excel.active
    sheet.title = "Assumptions"
    sheet.append(["Growth", 0.08, "=B1*100"])
    sheet["B1"].number_format = "0.0%"
    excel.create_sheet("Hidden").sheet_state = "hidden"
    excel.save(source)
    report.write_text(json.dumps({"value": 0.12}))
    with Workbook(source, tmp_path / "cache") as book:
        book.patch_cells(
            [{"operation": "SET_VALUE", "sheet": "Assumptions", "cell": "B1", "value": 0.12}],
            output,
        )
    baseline = cell_fingerprints(source, expected_changes("edit", 10))
    assert judge("edit", source, output, report, 10, baseline)["passed"]
    wrong = tmp_path / "wrong.xlsx"
    excel = load_workbook(output)
    excel["Assumptions"]["A1"] = "unrequested change"
    excel.save(wrong)
    verdict = judge("edit", source, wrong, report, 10, baseline)
    assert not verdict["passed"]
    assert verdict["checks"]["requested_values"]
    assert not verdict["checks"]["unchanged_values_formulas_structure_styles"]
    report.write_text(json.dumps({"dump": "x" * 6000}))
    assert not judge("edit", source, output, report, 10, baseline)["checks"]["bounded_json"]


def test_judge_reports_missing_output_instead_of_crashing(tmp_path):
    result = judge("create", None, tmp_path / "missing.xlsx", tmp_path / "missing.json", 10)
    assert not result["passed"] and "FileNotFoundError" in result["error"]
