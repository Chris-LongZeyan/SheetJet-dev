import os
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook as Excel
from openpyxl.worksheet.table import Table

from sheetjet import BudgetExceeded, SheetJetError, UnsupportedOperation, Workbook
from sheetjet.core import Budget
from sheetjet.index import dependencies


def test_missing_table_rows_are_null_rows(tmp_path, cache):
    path = tmp_path / "sparse.xlsx"
    b = Excel()
    s = b.active
    s.append(["ID", "Amount"])
    s["A2"] = 1
    s["B2"] = 10
    s["A5"] = 2
    s["B5"] = 20
    s.add_table(Table(displayName="Sparse", ref="A1:B7"))
    b.save(path)
    with Workbook(path, cache) as w:
        assert (
            w.query_engine.load("Sparse", schema={"ID": "BIGINT", "Amount": "BIGINT"})["rows"] == 6
        )
        assert w.query_engine.query("SELECT count(*),count(ID),sum(Amount) FROM Sparse")[
            "rows"
        ] == [[6, 2, 30]]


def test_null_clear_and_style_only_no_recalc(model, cache, tmp_path):
    out = tmp_path / "clear.xlsx"
    with Workbook(model, cache) as w:
        report = w.patch_cells(
            [{"operation": "SET_VALUE", "sheet": "Assumptions", "cell": "B1", "value": None}], out
        )
    with Workbook(out, cache) as w:
        assert w.read_range("Assumptions", "B1")["rows"] == [[None]]
    styled = tmp_path / "styled.xlsx"
    with Workbook(model, cache) as w:
        report = w.patch_cells(
            [
                {
                    "operation": "COPY_STYLE",
                    "sheet": "Assumptions",
                    "source_cell": "B1",
                    "cell": "B2",
                }
            ],
            styled,
        )
        assert not report["requires_recalculation"]
        assert report["changed_parts"] == ["xl/worksheets/sheet2.xml"]


def test_signed_package_and_unsupported_format(model, cache, tmp_path):
    with ZipFile(model, "a", compression=ZIP_DEFLATED) as z:
        z.writestr("_xmlsignatures/sig1.xml", b"<signature/>")
    with Workbook(model, cache) as w, pytest.raises(UnsupportedOperation, match="signatures"):
        w.patch_cells(
            [{"operation": "SET_VALUE", "sheet": "Assumptions", "cell": "B1", "value": 1}],
            tmp_path / "signed.xlsx",
        )
    bad = tmp_path / "book.xlsb"
    bad.write_bytes(b"unsupported")
    with pytest.raises(UnsupportedOperation, match="XLSB"):
        Workbook(bad, cache)


def test_context_long_text_and_sql_parameterization(model, cache):
    with Workbook(model, cache, Budget(max_chars=80)) as w, pytest.raises(BudgetExceeded):
        w.read_range("Assumptions", "A1:B2")
    with Workbook(model, cache) as w:
        e = w.query_engine
        e.load("Sales", columns=["Customer"], schema={"Customer": "VARCHAR"})
        assert e.query("SELECT count(*) FROM Sales WHERE Customer = ?", ["' OR 1=1 --"])[
            "rows"
        ] == [[0]]
        with pytest.raises(SheetJetError, match="one SELECT"):
            e.query("SELECT 1; SELECT 2")


def test_dynamic_dependencies_are_marked_incomplete():
    assert not dependencies('=INDIRECT("A1")')["complete"]
    assert not dependencies("=OFFSET(A1,1,1)")["complete"]
    assert dependencies("=SUM('Other Sheet'!A1:B2)")["references"] == ["'Other Sheet'!A1:B2"]


def test_blank_expected_value_and_range_scope(model, cache):
    with Workbook(model, cache) as w:
        assert w.validate_range("Assumptions", "D1:D2", {"D1": None})["valid"]
        with pytest.raises(SheetJetError, match="inside"):
            w.validate_range("Assumptions", "D1:D2", {"A1": None})


def test_label_scope_after_full_index(model, cache):
    with Workbook(model, cache) as w:
        assert w.find_text("Need full search", ["Assumptions"], full=True)["matches"]
        assert not w.find_text("Need full search", ["Assumptions"])["matches"]


def test_staging_cache_matches_sql_case_insensitive_names(model, cache):
    with Workbook(model, cache) as w:
        e = w.query_engine
        e.load(
            "Alias",
            sheet="Transactions",
            ref="A1:C3",
            columns=["Region"],
            schema={"Region": "VARCHAR"},
        )
        e.load(
            "alias",
            sheet="Transactions",
            ref="A1:C3",
            columns=["Revenue"],
            schema={"Revenue": "BIGINT"},
        )
        result = e.load(
            "Alias",
            sheet="Transactions",
            ref="A1:C3",
            columns=["Region"],
            schema={"Region": "VARCHAR"},
        )
        assert not result.get("cached", False)
        assert e.query("SELECT Region FROM Alias ORDER BY Region")["rows"] == [["APAC"], ["EMEA"]]


def test_index_rebuilds_changed_sheet_only(model, cache, tmp_path):
    with Workbook(model, cache) as w:
        w.index.build("Transactions")
        w.index.build("Assumptions")
        out = tmp_path / "changed.xlsx"
        w.patch_cells(
            [{"operation": "SET_VALUE", "sheet": "Assumptions", "cell": "B1", "value": 1}], out
        )
    # External replacement between sessions is supported; active-session mutation is not.
    os.replace(out, model)
    with Workbook(model, cache) as w:
        assert w.index.build("Transactions")["formula_count"] == 100
        assert w.metrics.counts["index_cache_hits"] == 1
        w.index.build("Assumptions")
        assert w.metrics.counts["index_cache_hits"] == 1
        assert w.metrics.counts["cells_scanned"] > 0
