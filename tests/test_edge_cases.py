from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from lxml import etree as ET
from openpyxl import Workbook as Excel
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table

from benchmarks.fixtures import suite
from sheetjet import UnsupportedOperation, Workbook
from sheetjet.core import NS, Q


def rewrite(path, changes, extra=None):
    temp = path.with_suffix(".tmp")
    with ZipFile(path) as src, ZipFile(temp, "w", compression=ZIP_DEFLATED) as dst:
        for info in src.infolist():
            raw = src.read(info.filename)
            if info.filename in changes:
                raw = changes[info.filename](raw)
            dst.writestr(info, raw)
        for name, raw in (extra or {}).items():
            dst.writestr(name, raw)
    temp.replace(path)


def test_shared_strings_rich_text_and_reuse(model, cache):
    def sheet(raw):
        root = ET.fromstring(raw)
        c = root.find(f'.//{Q}c[@r="A1"]')
        for child in list(c):
            c.remove(child)
        c.set("t", "s")
        ET.SubElement(c, Q + "v").text = "0"
        return ET.tostring(root)

    def rels(raw):
        root = ET.fromstring(raw)
        ET.SubElement(
            root,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            Id="rIdStrings",
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
            Target="sharedStrings.xml",
        )
        return ET.tostring(root)

    rewrite(
        model,
        {"xl/worksheets/sheet2.xml": sheet, "xl/_rels/workbook.xml.rels": rels},
        {
            "xl/sharedStrings.xml": f'<sst xmlns="{NS}" count="1" uniqueCount="1"><si><r><t>China </t></r><r><t>growth</t></r><rPh sb="0" eb="1"><t>NOT DISPLAYED</t></rPh></si></sst>'.encode()
        },
    )
    with Workbook(model, cache) as w:
        assert w.read_range("Assumptions", "A1")["rows"] == [["China growth"]]
        assert w.metrics.counts["shared_strings_loaded"] == 1
    with Workbook(model, cache) as w:
        assert w.read_range("Assumptions", "A1")["rows"] == [["China growth"]]
        assert w.metrics.counts.get("shared_strings_loaded", 0) == 0


def test_shared_formula_group_refuses_follower_patch(model, cache, tmp_path):
    def change(raw):
        root = ET.fromstring(raw)
        a = root.find(f'.//{Q}c[@r="E2"]/{Q}f')
        a.attrib.update({"t": "shared", "si": "0", "ref": "E2:E101"})
        for row in range(3, 102):
            f = root.find(f'.//{Q}c[@r="E{row}"]/{Q}f')
            f.attrib.update({"t": "shared", "si": "0"})
            f.text = None
        return ET.tostring(root)

    rewrite(model, {"xl/worksheets/sheet1.xml": change})
    with Workbook(model, cache) as w:
        assert w.formula_patterns("Transactions")[0]["count"] == 100
        with pytest.raises(UnsupportedOperation, match="formula group"):
            w.patch_cells(
                [{"operation": "SET_VALUE", "sheet": "Transactions", "cell": "E90", "value": 1}],
                tmp_path / "bad.xlsx",
            )


def test_prefixed_xml_and_self_closing_rows(model, cache, tmp_path):
    raw = f'<m:worksheet xmlns:m="{NS}"><m:dimension ref="A1:C5"/><m:sheetData><m:row r="1"><m:c r="A1" t="n"><m:v>1</m:v></m:c><m:c r="C1"/></m:row><m:row r="2"/><m:row r="5"><m:c r="A5" t="n"><m:v>5</m:v></m:c></m:row></m:sheetData></m:worksheet>'.encode()
    rewrite(model, {"xl/worksheets/sheet2.xml": lambda _: raw})
    out = tmp_path / "prefix.xlsx"
    with Workbook(model, cache) as w:
        w.write_range("Assumptions", "B1", [[2, 3], [4, 5], [6, 7]], out)
    with Workbook(out, cache) as w:
        assert w.read_range("Assumptions", "A1:C3")["rows"] == [
            [1, 2, 3],
            [None, 4, 5],
            [None, 6, 7],
        ]


def test_calculation_chain_is_removed(model, cache, tmp_path):
    def rels(raw):
        root = ET.fromstring(raw)
        ET.SubElement(
            root,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            Id="rIdChain",
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain",
            Target="calcChain.xml",
        )
        return ET.tostring(root)

    def types(raw):
        root = ET.fromstring(raw)
        ET.SubElement(
            root,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            PartName="/xl/calcChain.xml",
            ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml",
        )
        return ET.tostring(root)

    rewrite(
        model,
        {"xl/_rels/workbook.xml.rels": rels, "[Content_Types].xml": types},
        {"xl/calcChain.xml": f'<calcChain xmlns="{NS}"><c r="E2" i="1"/></calcChain>'.encode()},
    )
    with Workbook(model, cache) as w:
        out = tmp_path / "chain.xlsx"
        result = w.patch_cells(
            [
                {
                    "operation": "SET_FORMULA",
                    "sheet": "Transactions",
                    "cell": "E2",
                    "formula": "=C2*2",
                }
            ],
            out,
        )
        assert result["removed_parts"] == ["xl/calcChain.xml"]
    b = load_workbook(out)
    assert b.calculation.forceFullCalc
    b.close()


def test_new_calc_properties_precede_extension_list(model, cache, tmp_path):
    def change(raw):
        root = ET.fromstring(raw)
        calc = root.find(Q + "calcPr")
        root.remove(calc)
        ET.SubElement(root, Q + "extLst")
        return ET.tostring(root)

    rewrite(model, {"xl/workbook.xml": change})
    out = tmp_path / "calc-order.xlsx"
    with Workbook(model, cache) as w:
        w.patch_cells(
            [{"operation": "SET_VALUE", "sheet": "Assumptions", "cell": "B1", "value": 0.1}], out
        )
    with ZipFile(out) as z:
        tags = [n.tag for n in ET.fromstring(z.read("xl/workbook.xml"))]
        assert tags.index(Q + "calcPr") < tags.index(Q + "extLst")


def test_fixture_suite_preserves_pivot_and_macro_payloads(tmp_path, cache):
    fixtures = tmp_path / "fixtures"
    manifest = suite(fixtures, 20)
    assert len(manifest) == 10
    for filename in ["charts-pivots.xlsx", "macro-preservation.xlsm"]:
        source = fixtures / filename
        output = tmp_path / filename
        with Workbook(source, cache) as w:
            meta = w.inspect_workbook()
            if filename.endswith(".xlsx"):
                assert meta["features"]["pivots"] == 1
            else:
                assert meta["features"]["vba"]
            report = w.patch_cells(
                [{"operation": "SET_VALUE", "sheet": "Model", "cell": "B2", "value": 123}], output
            )
        with ZipFile(source) as a, ZipFile(output) as b:
            for name in a.namelist():
                if name not in report["changed_parts"]:
                    assert a.read(name) == b.read(name)
        if filename.endswith(".xlsx"):
            b = load_workbook(output)
            assert len(b["Model"]._pivots) == 1
            b.close()


def test_typed_table_empty_text_null_and_decimal(tmp_path, cache):
    path = tmp_path / "typed.xlsx"
    b = Excel()
    s = b.active
    s.append(["Key", "Amount"])
    s.append(["0001", 1.25])
    s.append(["", None])
    s.append([None, 2.50])
    s.add_table(Table(displayName="Typed", ref="A1:B4"))
    b.save(path)
    with Workbook(path, cache) as w:
        w.query_engine.load("Typed", schema={"Amount": "DECIMAL(12,2)", "Key": "VARCHAR"})
        assert w.query_engine.query(
            'SELECT "Key","Amount" FROM Typed ORDER BY "Amount" NULLS LAST'
        )["rows"] == [["0001", "1.25"], [None, "2.50"], ["", None]]
        result = w.profile_table("Typed", {"Key": "VARCHAR", "Amount": "DECIMAL(12,2)"})
        assert result["rows"][1][-1] == 1.875


def test_query_reconciliation_between_distinct_tables(tmp_path, cache):
    path = tmp_path / "join.xlsx"
    b = Excel()
    s = b.active
    for row in [
        ["ID", "Sales", None, "ID", "Cost"],
        ["A", 100, None, "A", 60],
        ["B", 200, None, "B", 150],
    ]:
        s.append(row)
    b.save(path)
    with Workbook(path, cache) as w:
        e = w.query_engine
        e.load("Sales", sheet="Sheet", ref="A1:B3", schema={"ID": "VARCHAR", "Sales": "BIGINT"})
        e.load("Costs", sheet="Sheet", ref="D1:E3", schema={"ID": "VARCHAR", "Cost": "BIGINT"})
        assert e.query(
            "SELECT a.ID, a.Sales-b.Cost AS margin FROM Sales a JOIN Costs b USING(ID) ORDER BY 1"
        )["rows"] == [["A", 40], ["B", 50]]
