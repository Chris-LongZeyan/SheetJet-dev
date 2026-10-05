from __future__ import annotations

import copy
import os
from collections import OrderedDict
from pathlib import Path

from .core import Budget, Metrics, Q, address, bounds, compact, position
from .errors import SheetJetError
from .index import Index, dependencies
from .package import Package, xml


class Workbook:
    """A session owns its open ZIP, index and small range cache. Use as a context manager."""

    address = staticmethod(address)

    def __init__(self, path, cache_dir=None, budget=None):
        self.budget = budget or Budget()
        self.metrics = Metrics()
        cache_dir = (
            cache_dir or Path(os.environ.get("LOCALAPPDATA", Path.home() / ".cache")) / "sheetjet"
        )
        with self.metrics.time("open"):
            self.package = Package(path, cache_dir, self.metrics)
        self._index = None
        self._ranges = OrderedDict()
        self._query = None

    @property
    def index(self):
        if self._index is None:
            self.package.check()
            self._index = Index(self.package)
        return self._index

    def serialize(self, value, budget=None):
        return compact(value, budget or self.budget, self.metrics)

    def inspect_workbook(self, offset=0, limit=20):
        self.package.check()
        if offset < 0 or not 1 <= limit <= 100:
            raise SheetJetError("Use nonnegative offset and a limit between 1 and 100")
        p = self.package
        unpacked = sum(i.file_size for i in p.parts.values())
        strategy = (
            "huge"
            if unpacked > 1024**3
            else "large"
            if unpacked > 100 * 1024**2
            else "medium"
            if unpacked > 10 * 1024**2
            else "small"
        )
        sheets = list(p.sheets.values())
        result = {
            "file": p.path.name,
            "bytes": p.path.stat().st_size,
            "xml_bytes": unpacked,
            "strategy": strategy,
            "sheet_count": len(sheets),
            "sheets": sheets[offset : offset + limit],
            "next_offset": offset + limit if offset + limit < len(sheets) else None,
            "named_range_count": len(p.names),
            "features": p.features,
            "calculation": p.calculation,
            "counts_note": "Formula/nonempty counts and accurate dimensions require inspect_sheet(deep=True)",
        }
        self.serialize(result)
        return result

    def inspect_sheet(self, sheet, deep=False, section="summary", offset=0, limit=20):
        if offset < 0 or not 1 <= limit <= 100:
            raise SheetJetError("Use nonnegative offset and a limit between 1 and 100")
        self.package.check()
        if sheet not in self.package.sheets:
            raise SheetJetError(f"Unknown sheet {sheet!r}")
        meta = self.index.build(sheet) if deep else self.package.sheets[sheet]
        if section == "summary":
            result = {
                k: v
                for k, v in meta.items()
                if k
                not in {
                    "merged_ranges",
                    "hidden_rows",
                    "hidden_columns",
                    "formula_groups",
                    "style_counts",
                }
            }
            for key in (
                "merged_ranges",
                "hidden_rows",
                "hidden_columns",
                "formula_groups",
                "style_counts",
            ):
                if key in meta:
                    result[key + "_count"] = len(meta[key])
        else:
            if not deep or section not in meta:
                raise SheetJetError(
                    "Detailed sections require deep=True and a valid metadata section"
                )
            data = meta[section]
            data = list(data.items()) if isinstance(data, dict) else data
            if not isinstance(data, list):
                raise SheetJetError("This section is not pageable")
            result = {
                "sheet": sheet,
                "section": section,
                "items": data[offset : offset + limit],
                "total": len(data),
            }
        self.serialize(result)
        return result

    def find_text(self, query, sheets=None, limit=20, full=False):
        matches = self.index.search(query, sheets, limit, "full" if full else "labels")
        result = {
            "matches": matches,
            "scope": "all text, first 2048 characters per cell"
            if full
            else "first 20 rows and first 3 columns; use full=True for remaining text",
        }
        self.serialize(result)
        return result

    def find_table(self, query):
        self.package.check()
        result = [
            t for name, t in self.package.tables.items() if query.casefold() in name.casefold()
        ]
        self.serialize(result)
        return result

    def find_names(self, query="", limit=20):
        self.package.check()
        result = [
            n for n in self.package.names if query.casefold() in n.get("name", "").casefold()
        ][:limit]
        self.serialize(result)
        return result

    def get_cells(self, sheet, ref):
        self.package.check()
        self.budget.check_range(ref)
        key = (sheet, ref)
        if key in self._ranges:
            self.metrics.add("range_cache_hits")
            self._ranges.move_to_end(key)
            return copy.deepcopy(self._ranges[key])
        with self.metrics.time("range_read"):
            cells = [cell for _, row in self.package.rows(sheet, ref) for cell in row]
        self.metrics.add("range_reads")
        # Oversized cell text may be useful to local callers, but must not fill the cache.
        if (
            sum(len(str(c["value"])) + len(c.get("formula", "")) for c in cells)
            <= self.budget.max_chars
        ):
            self._ranges[key] = copy.deepcopy(cells)
        if len(self._ranges) > 16:
            self._ranges.popitem(last=False)
        return cells

    def read_cells(self, sheet, cells):
        """Internal sparse batch read: one sequential scan, even for distant coordinates."""
        wanted = {address(*position(c)) for c in cells}
        if len(wanted) > 10000:
            raise SheetJetError("Sparse read exceeds 10000-cell execution budget")
        if not wanted:
            return []
        result = {c: {"cell": c, "value": None, "type": "n", "style": 0} for c in wanted}
        maxrow = max(position(c)[0] for c in wanted)
        # Decode only requested cells; shared string tables remain unopened for numeric-only reads.
        from lxml import etree as ET

        from .package import release

        self.package.check()
        with self.package.zip.open(self.package.sheets[sheet]["part"]) as stream:
            for _, row in ET.iterparse(
                stream, events=("end",), tag=Q + "row", resolve_entities=False, no_network=True
            ):
                if int(row.get("r")) > maxrow:
                    break
                for node in row.iterchildren(Q + "c"):
                    if node.get("r") in wanted:
                        result[node.get("r")] = self.package.cell(node)
                release(row)
        self.metrics.add("range_reads")
        return [result[address(*position(c))] for c in cells]

    def read_range(self, sheet, ref):
        cells = self.get_cells(sheet, ref)
        r1, c1, r2, c2 = bounds(ref)
        rows = [[None] * (c2 - c1 + 1) for _ in range(r2 - r1 + 1)]
        formulas, styles = {}, {}
        for cell in cells:
            r, c = position(cell["cell"])
            rows[r - r1][c - c1] = cell["value"]
            if "formula" in cell:
                formulas[cell["cell"]] = {"text": cell["formula"], **cell["formula_attributes"]}
            if cell["style"]:
                styles[cell["cell"]] = cell["style"]
        result = {"sheet": sheet, "range": ref, "rows": rows}
        if formulas:
            result.update(formulas=formulas, formula_values="stored caches; may be absent or stale")
        if styles:
            result["styles"] = styles
        result["dates"] = (
            "Excel serial numbers retain their style IDs; no implicit timezone/date conversion"
        )
        self.serialize(result)
        self.metrics.add("result_cells", (r2 - r1 + 1) * (c2 - c1 + 1))
        return result

    def get_formulas(self, sheet, ref):
        result = [
            {"cell": c["cell"], "formula": c["formula"], "attributes": c["formula_attributes"]}
            for c in self.get_cells(sheet, ref)
            if "formula" in c
        ]
        self.serialize(result)
        return result

    def get_dependencies(self, sheet, ref):
        result = [
            {
                "cell": c["cell"],
                **(
                    dependencies(c["formula"])
                    if c["formula"] != "="
                    else {
                        "references": [],
                        "complete": False,
                        "scope": "Shared/data-table formula follower; inspect its anchor",
                    }
                ),
            }
            for c in self.get_formulas(sheet, ref)
        ]
        self.serialize(result)
        return result

    def formula_patterns(self, sheet, limit=20, offset=0):
        result = self.index.patterns(sheet, limit, offset)
        self.serialize(result)
        return result

    def get_styles(self, sheet, ref):
        cells = self.get_cells(sheet, ref)
        root = xml(self.package.zip.read("xl/styles.xml"))
        xfs = root.find(Q + "cellXfs")
        ids = sorted({c["style"] for c in cells})
        formats = {
            int(n.get("numFmtId")): n.get("formatCode")
            for n in root.findall(Q + "numFmts/" + Q + "numFmt")
        }
        from lxml import etree as ET
        from openpyxl.styles.numbers import BUILTIN_FORMATS

        defs = {}
        for i in ids:
            xf = xfs[i]
            data = dict(xf.attrib)
            fid = int(xf.get("numFmtId", "0"))
            data["number_format"] = formats.get(fid, BUILTIN_FORMATS.get(fid))
            for tag in ("alignment", "protection"):
                child = xf.find(Q + tag)
                if child is not None:
                    data[tag] = dict(child.attrib)
            for attr, group in (("fontId", "fonts"), ("fillId", "fills"), ("borderId", "borders")):
                nodes = root.find(Q + group)
                if nodes is not None:
                    data[group] = ET.tostring(nodes[int(xf.get(attr, "0"))], encoding="unicode")
            defs[i] = data
        result = {"cells": {c["cell"]: c["style"] for c in cells}, "styles": defs}
        self.serialize(result)
        return result

    @property
    def query_engine(self):
        if self._query is None:
            from .query import QueryEngine

            self._query = QueryEngine(self)
        return self._query

    def read_table(
        self, name, columns=None, filters=None, limit=20, schema=None, allow_cached_formulas=False
    ):
        from .query import identifier

        engine = self.query_engine
        engine.load(
            name, schema=schema, columns=columns, allow_cached_formulas=allow_cached_formulas
        )
        conditions, params = [], []
        available = columns or self.package.tables[name]["columns"]
        for col, value in (filters or {}).items():
            if col not in available:
                raise SheetJetError("Filter column must be included in the projection")
            if value is None:
                conditions.append(identifier(col) + " IS NULL")
            else:
                conditions.append(identifier(col) + " = ?")
                params.append(value)
        if not 1 <= limit <= self.budget.max_rows:
            raise SheetJetError("Limit must fit the row budget")
        sql = (
            "SELECT * FROM "
            + identifier(name)
            + (" WHERE " + " AND ".join(conditions) if conditions else "")
            + f" LIMIT {limit}"
        )
        return engine.query(sql, params)

    def profile_table(self, name, schema):
        from .query import identifier

        self.query_engine.load(name, schema=schema, columns=list(schema))
        # A long-form profile stays bounded even for tables with many rows.
        terms = []
        for col, kind in schema.items():
            label = "'" + col.replace("'", "''") + "'"
            c = identifier(col)
            numeric = kind.upper().startswith(("DOUBLE", "BIGINT", "DECIMAL"))
            terms.append(
                f"SELECT {label} AS column_name, count(*) AS rows, count({c}) AS nonnull, count(DISTINCT {c}) AS distinct_values, "
                + (f"avg({c})" if numeric else "NULL")
                + f" AS mean FROM {identifier(name)}"
            )
        return self.query_engine.query(" UNION ALL ".join(terms))

    def patch_cells(self, operations, output, overwrite=False):
        from .patch import apply_patch

        return apply_patch(self, operations, output, overwrite)

    def write_range(self, sheet, start, rows, output):
        r, c = position(start)
        if not rows or len({len(row) for row in rows}) != 1:
            raise SheetJetError("Rows must form a nonempty rectangle")
        ops = [
            {"operation": "SET_VALUE", "sheet": sheet, "cell": address(r + i, c + j), "value": v}
            for i, row in enumerate(rows)
            for j, v in enumerate(row)
        ]
        return self.patch_cells(ops, output)

    def validate_range(self, sheet, ref, expected=None):
        from .validation import validate_range

        return validate_range(self, sheet, ref, expected)

    def close(self):
        if self._query:
            self._query.close()
        if self._index:
            self._index.close()
        self.package.close()

    def __enter__(self):
        return self

    def __exit__(self, *_):
        self.close()
