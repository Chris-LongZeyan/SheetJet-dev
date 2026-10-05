"""Every measurement runs in a fresh process; no API keys or LLM calls required."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
import statistics
import subprocess
import sys
import tempfile
import threading
import time
from collections import defaultdict
from pathlib import Path


def worker(method, path, rows):
    import psutil

    process = psutil.Process()
    stop = threading.Event()
    peak = [process.memory_info().rss]

    def sample():
        while not stop.wait(0.01):
            peak[0] = max(peak[0], process.memory_info().rss)

    thread = threading.Thread(target=sample, daemon=True)
    thread.start()
    # Imports are included in timing for all methods; report this explicitly.
    start = time.perf_counter()
    exposed_chars = 0
    metrics = {}
    with tempfile.TemporaryDirectory(prefix="sheetjet-bench-") as tmp:
        if method.startswith("sheetjet"):
            from sheetjet import Workbook

            with Workbook(path, Path(tmp) / "cache") as w:
                if method == "sheetjet_inspect":
                    answer = w.inspect_workbook()
                elif method == "sheetjet_locate":
                    answer = w.read_range("Transactions", f"D{rows + 1}")["rows"][0][0]
                elif method in {"sheetjet_aggregate", "sheetjet_warm_query"}:
                    w.query_engine.load(
                        "Transactions",
                        columns=["Region", "Revenue"],
                        schema={"Region": "VARCHAR", "Revenue": "BIGINT"},
                    )
                    sql = 'SELECT "Region",sum("Revenue") AS revenue FROM Transactions GROUP BY 1 ORDER BY 1'
                    if method == "sheetjet_warm_query":
                        w.query_engine.query(sql)
                        start = time.perf_counter()
                    answer = w.query_engine.query(sql)["rows"]
                elif method == "sheetjet_patch":
                    answer = w.patch_cells(
                        [
                            {
                                "operation": "SET_VALUE",
                                "sheet": "Assumptions",
                                "cell": "B1",
                                "value": 0.12,
                            }
                        ],
                        Path(tmp) / "edited.xlsx",
                    )
                    answer = {
                        "changed_cells": answer["changed_cells"],
                        "untouched_parts_verified": answer["validation"][
                            "untouched_parts_crc_verified"
                        ],
                    }
                else:
                    raise ValueError(method)
                metrics = w.metrics.snapshot()
        else:
            from openpyxl import load_workbook

            book = load_workbook(path, read_only=True, data_only=True)
            sheet = book["Transactions"]
            if method == "openpyxl_inspect":
                answer = {"sheets": book.sheetnames, "range": sheet.calculate_dimension()}
            elif method == "openpyxl_locate":
                answer = sheet.cell(rows + 1, 4).value
            else:
                totals = defaultdict(int)
                # Naive baseline materializes full rows and their JSON, as an LLM-context workflow would.
                data = (
                    list(sheet.iter_rows(min_row=2, values_only=True))
                    if method == "naive_context"
                    else sheet.iter_rows(min_row=2, values_only=True)
                )
                if method == "naive_context":
                    exposed_chars = len(json.dumps(data, separators=(",", ":")))
                for row in data:
                    totals[row[2]] += row[3]
                answer = [[k, v] for k, v in sorted(totals.items())]
            book.close()
    elapsed = time.perf_counter() - start
    exposed_chars = exposed_chars or len(json.dumps(answer, separators=(",", ":")))
    peak[0] = max(peak[0], process.memory_info().rss)
    stop.set()
    thread.join()
    return {
        "method": method,
        "seconds": elapsed,
        "peak_rss_mb": peak[0] / 1024**2,
        "context_characters": exposed_chars,
        "estimated_context_tokens": math.ceil(exposed_chars / 4),
        "answer": answer,
        "metrics": metrics,
    }


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--rows", type=int, default=10000)
    p.add_argument("--repeats", type=int, default=3)
    p.add_argument("--output", type=Path, default=Path("benchmark-output/results.json"))
    p.add_argument("--workbook", type=Path)
    p.add_argument(
        "--worker",
        choices=[
            "sheetjet_inspect",
            "sheetjet_locate",
            "sheetjet_aggregate",
            "sheetjet_warm_query",
            "sheetjet_patch",
            "openpyxl_inspect",
            "openpyxl_locate",
            "openpyxl_stream",
            "naive_context",
        ],
    )
    args = p.parse_args()
    if args.worker:
        print(json.dumps(worker(args.worker, args.workbook, args.rows), separators=(",", ":")))
        return
    args.output.parent.mkdir(parents=True, exist_ok=True)
    path = args.workbook or args.output.parent / "transactions.xlsx"
    if args.workbook is None:
        from .fixtures import transactions

        transactions(path, args.rows)
    methods = [
        "sheetjet_inspect",
        "openpyxl_inspect",
        "sheetjet_locate",
        "openpyxl_locate",
        "sheetjet_aggregate",
        "openpyxl_stream",
        "naive_context",
        "sheetjet_warm_query",
        "sheetjet_patch",
    ]
    measurements = []
    for method in methods:
        for repetition in range(args.repeats):
            result = subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "benchmarks.run",
                    "--worker",
                    method,
                    "--rows",
                    str(args.rows),
                    "--workbook",
                    str(path),
                ],
                capture_output=True,
                text=True,
                check=True,
            )
            m = json.loads(result.stdout)
            measurements.append(m)
            print(
                f"{method} {repetition + 1}/{args.repeats}: {m['seconds']:.3f}s, {m['peak_rss_mb']:.1f} MiB",
                flush=True,
            )
    regions = ["APAC", "EMEA", "Americas", "Other"]
    expected = defaultdict(int)
    for i in range(1, args.rows + 1):
        expected[regions[i % 4]] += i % 1000 + 1
    expected = [[k, v] for k, v in sorted(expected.items())]
    for m in measurements:
        if m["method"] in {
            "sheetjet_aggregate",
            "sheetjet_warm_query",
            "openpyxl_stream",
            "naive_context",
        }:
            assert m["answer"] == expected, (m["method"], m["answer"], expected)
        if "locate" in m["method"]:
            assert m["answer"] == args.rows % 1000 + 1
    medians = [
        {
            "method": method,
            **{
                key: statistics.median(m[key] for m in measurements if m["method"] == method)
                for key in [
                    "seconds",
                    "peak_rss_mb",
                    "context_characters",
                    "estimated_context_tokens",
                ]
            },
        }
        for method in methods
    ]
    import duckdb
    import lxml
    import openpyxl

    report = {
        "environment": {
            "python": platform.python_version(),
            "platform": platform.platform(),
            "cpu": platform.processor(),
            "duckdb": duckdb.__version__,
            "openpyxl": openpyxl.__version__,
            "lxml": lxml.__version__,
        },
        "rows": args.rows,
        "cells": (args.rows + 1) * 5 + 2,
        "repeats": args.repeats,
        "file_bytes": path.stat().st_size,
        "file_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "correctness": "all aggregation and lookup answers matched independent fixture arithmetic",
        "methodology": "Fresh process per trial; OS file cache not flushed. Wall time includes imports, open, staging, execution, cleanup except warm_query (only repeated query and cleanup). Peak RSS sampled every 10ms includes imports/staging. Token estimate = ceil(JSON characters/4), not a tokenizer count or measured LLM usage. Context includes full rows only for naive_context, bounded answer for other methods. Synthetic data, not universal performance rankings.",
        "medians": medians,
        "measurements": measurements,
    }
    args.output.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(f"Saved {args.output}")


if __name__ == "__main__":
    main()
