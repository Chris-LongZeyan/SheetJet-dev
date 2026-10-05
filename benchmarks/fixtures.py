from __future__ import annotations

import argparse
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table


def transactions(path, rows=1000000):
    """Write a standards-based, inline-string transaction fixture without a cell object graph."""
    if not 1 <= rows <= 1048575:
        raise ValueError("Rows must fit one Excel worksheet including its header")
    book = Workbook()
    s = book.active
    s.title = "Transactions"
    s.append(["ID", "Customer", "Region", "Revenue", "Cost"])
    s.append([1, "Customer1", "APAC", 2, 1])
    s.add_table(Table(displayName="Transactions", ref=f"A1:E{rows + 1}"))
    s.freeze_panes = "D2"
    a = book.create_sheet("Assumptions")
    a.append(["China revenue growth assumption", 0.08])
    a["B1"].number_format = "0.0%"
    template = path.with_suffix(".template.xlsx")
    book.save(template)
    try:
        with (
            ZipFile(template) as source,
            ZipFile(path, "w", compression=ZIP_DEFLATED, compresslevel=1) as out,
        ):
            for info in source.infolist():
                if info.filename != "xl/worksheets/sheet1.xml":
                    out.writestr(info, source.read(info.filename))
                    continue
                with out.open(info.filename, "w", force_zip64=True) as stream:
                    stream.write(
                        (
                            f'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><dimension ref="A1:E{rows + 1}"/><sheetViews><sheetView workbookViewId="0"><pane xSplit="3" ySplit="1" topLeftCell="D2" state="frozen"/></sheetView></sheetViews><sheetData><row r="1">'
                        ).encode()
                    )
                    for col, text in zip(
                        "ABCDE", ["ID", "Customer", "Region", "Revenue", "Cost"], strict=True
                    ):
                        stream.write(
                            f'<c r="{col}1" t="inlineStr"><is><t>{text}</t></is></c>'.encode()
                        )
                    stream.write(b"</row>")
                    regions = ["APAC", "EMEA", "Americas", "Other"]
                    for i in range(1, rows + 1):
                        r = i + 1
                        stream.write(
                            (
                                f'<row r="{r}"><c r="A{r}" t="n"><v>{i}</v></c><c r="B{r}" t="inlineStr"><is><t>Customer{i % 1000}</t></is></c><c r="C{r}" t="inlineStr"><is><t>{regions[i % 4]}</t></is></c><c r="D{r}" t="n"><v>{i % 1000 + 1}</v></c><c r="E{r}" t="n"><v>{i % 500}</v></c></row>'
                            ).encode()
                        )
                    stream.write(
                        b'</sheetData><tableParts count="1"><tablePart r:id="rId1"/></tableParts></worksheet>'
                    )
    finally:
        template.unlink(missing_ok=True)
    return path


def suite(directory, rows=10000):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    manifest = []
    path = transactions(directory / "transactions.xlsx", rows)
    manifest.append({"file": path.name, "kind": "transactions", "rows": rows, "native": True})
    kinds = [
        "financial-model",
        "formula-heavy",
        "formatting-heavy",
        "hidden-names",
        "merged",
        "charts-pivots",
        "messy",
        "interconnected",
        "macro-preservation",
    ]
    for kind in kinds:
        book = Workbook()
        s = book.active
        s.title = "Model"
        s.append(["Metric", "Value", "Forecast"])
        if kind in {"financial-model", "formula-heavy"}:
            n = min(rows, 50000) if kind == "formula-heavy" else 200
            for r in range(2, n + 2):
                s.append([f"Metric {r}", r, f"=B{r}*1.08"])
            a = book.create_sheet("Assumptions")
            a.append(["Growth", 0.08])
            s["C2"] = "=B2*(1+Assumptions!B1)"
        elif kind == "formatting-heavy":
            for r in range(2, min(rows, 10000) + 2):
                s.append([f"Label {r}", r, r / 100])
                for c in s[r]:
                    c.fill = PatternFill("solid", fgColor=f"FF{r % 256:02X}80A0")
                    c.font = Font(name="Calibri", bold=r % 2 == 0, size=10 + r % 3)
                    c.number_format = "0.00%" if c.column == 3 else "#,##0"
        elif kind == "hidden-names":
            for i in range(10):
                a = book.create_sheet(f"Hidden{i}")
                a.sheet_state = "veryHidden" if i % 2 else "hidden"
                a["B2"] = i
                book.defined_names.add(DefinedName(f"Input{i}", attr_text=f"'Hidden{i}'!$B$2"))
            s.row_dimensions[5].hidden = True
            s.column_dimensions["C"].hidden = True
        elif kind == "merged":
            for r in range(3, 200, 3):
                s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
                s.cell(r, 1, f"Section {r}")
        elif kind == "charts-pivots":
            for r in range(2, 30):
                s.append([f"Region {r % 4}", r, r * 2])
            chart = BarChart()
            chart.add_data(
                Reference(s, min_col=2, max_col=3, min_row=1, max_row=20), titles_from_data=True
            )
            s.add_chart(chart, "F2")
            # Native pivot metadata/cache generated through openpyxl, with no external inputs.
            from openpyxl.pivot.cache import (
                CacheDefinition,
                CacheField,
                CacheSource,
                SharedItems,
                WorksheetSource,
            )
            from openpyxl.pivot.fields import Text
            from openpyxl.pivot.table import (
                DataField,
                Location,
                PivotField,
                RowColField,
                TableDefinition,
            )

            fields = [
                CacheField(
                    name="Metric",
                    sharedItems=SharedItems(_fields=[Text(v=f"Region {i}") for i in range(4)]),
                ),
                CacheField(
                    name="Value", sharedItems=SharedItems(containsNumber=True, containsString=False)
                ),
                CacheField(
                    name="Forecast",
                    sharedItems=SharedItems(containsNumber=True, containsString=False),
                ),
            ]
            cache = CacheDefinition(
                cacheSource=CacheSource(
                    type="worksheet", worksheetSource=WorksheetSource(ref="A1:C29", sheet="Model")
                ),
                cacheFields=fields,
                refreshOnLoad=True,
                recordCount=0,
            )
            pivot = TableDefinition(
                name="RegionTotals",
                cacheId=1,
                dataCaption="Values",
                location=Location(ref="J1:K6", firstHeaderRow=1, firstDataRow=1, firstDataCol=1),
                pivotFields=[PivotField(axis="axisRow"), PivotField(dataField=True), PivotField()],
                rowFields=[RowColField(x=0)],
                dataFields=[DataField(name="Sum of Value", fld=1, subtotal="sum")],
            )
            pivot.cache = cache
            s.add_pivot(pivot)
        elif kind == "messy":
            s["A1"] = "Analyst working area"
            s.append(["Customer", "Sales", None, None, "Customer", "Costs"])
            for r in range(3, 100):
                s.append([f"C{r % 7}", r * 10, None, None, f"C{r % 7}", r * 6])
            s.add_table(Table(displayName="Sales", ref="A2:B99"))
            s.add_table(Table(displayName="Costs", ref="E2:F99"))
            s.merge_cells("A102:F102")
            s["A102"] = "Notes: units differ from the other workbook."
        elif kind == "interconnected":
            s["B2"] = 100
            for i in range(1, 50):
                a = book.create_sheet(f"Schedule{i}")
                previous = "Model" if i == 1 else f"Schedule{i - 1}"
                a["B2"] = f"='{previous}'!B2*1.01"
        else:
            s["B2"] = 100
        suffix = ".xlsm" if kind == "macro-preservation" else ".xlsx"
        path = directory / (kind + suffix)
        book.save(path)
        if kind == "macro-preservation":
            # Opaque canary tests byte retention only. It is deliberately NOT runnable VBA.
            from lxml import etree as ET

            pending = path.with_suffix(".pending")
            with ZipFile(path) as src, ZipFile(pending, "w", compression=ZIP_DEFLATED) as out:
                for info in src.infolist():
                    raw = src.read(info.filename)
                    if info.filename == "[Content_Types].xml":
                        root = ET.fromstring(raw)
                        for el in root:
                            if el.get("PartName") == "/xl/workbook.xml":
                                el.set(
                                    "ContentType",
                                    "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                                )
                        ET.SubElement(
                            root,
                            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                            PartName="/xl/vbaProject.bin",
                            ContentType="application/vnd.ms-office.vbaProject",
                        )
                        raw = ET.tostring(root)
                    elif info.filename == "xl/_rels/workbook.xml.rels":
                        root = ET.fromstring(raw)
                        ET.SubElement(
                            root,
                            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                            Id="rIdVBA",
                            Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject",
                            Target="vbaProject.bin",
                        )
                        raw = ET.tostring(root)
                    out.writestr(info, raw)
                out.writestr(
                    "xl/vbaProject.bin", b"SHEETJET OPAQUE VBA PRESERVATION CANARY\x00\xff" * 100
                )
            pending.replace(path)
        manifest.append(
            {
                "file": path.name,
                "kind": kind,
                "native": kind != "macro-preservation",
                "note": "Opaque canary, not executable VBA; supply a real XLSM for native macro verification"
                if kind == "macro-preservation"
                else "Native Excel rendering/recalculation not verified",
            }
        )
    (directory / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("directory", type=Path)
    p.add_argument("--rows", type=int, default=10000)
    p.add_argument("--transactions-only", action="store_true")
    a = p.parse_args()
    a.directory.mkdir(parents=True, exist_ok=True)
    if a.transactions_only:
        print(transactions(a.directory / "transactions.xlsx", a.rows))
    else:
        print(json.dumps(suite(a.directory, a.rows), indent=2))
