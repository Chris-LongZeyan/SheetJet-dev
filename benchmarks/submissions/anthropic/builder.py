"""Original standalone XLSX task builders; preserve untouched package XML."""
import json
import math
import posixpath
import re
import sys
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def sheet_paths(archive):
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {r.attrib["Id"]: r.attrib["Target"] for r in relationships}
    result = {}
    for sheet in workbook.find("{%s}sheets" % MAIN):
        target = targets[sheet.attrib["{%s}id" % REL]]
        result[sheet.attrib["name"]] = (target.lstrip("/") if target.startswith("/")
                                          else posixpath.normpath(posixpath.join("xl", target)))
    return result


def patch_numeric(xml, coordinate, value):
    """Replace only an existing numeric cell's value; preserve every other byte."""
    if not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError("Expected finite numeric value")
    cell = re.compile(rb'<c\b(?=[^>]*\br="' + coordinate.encode("ascii") + rb'")[^>]*>.*?</c>', re.DOTALL)
    matches = list(cell.finditer(xml))
    if len(matches) != 1:
        raise ValueError("Expected exactly one populated target cell: " + coordinate)
    match = matches[0]
    original = match.group()
    if re.search(rb'<f\b', original) or re.search(rb'\bt="(?!n")[^"]+"', original.split(b">", 1)[0]):
        raise ValueError("Refusing to replace nonnumeric cell " + coordinate)
    value_xml = ("<v>" + repr(value) + "</v>").encode("ascii")
    changed, count = re.subn(rb'<v\b[^>]*>.*?</v>', value_xml, original, flags=re.DOTALL)
    if count != 1:
        raise ValueError("Expected one numeric value in " + coordinate)
    return xml[:match.start()] + changed + xml[match.end():]


def save_patches(source, output, updates):
    if Path(source).resolve() == Path(output).resolve():
        raise ValueError("Output must differ from input")
    with ZipFile(source) as src:
        paths = sheet_paths(src)
        replacements = {}
        for sheet, cells in updates.items():
            part = paths[sheet]
            xml = src.read(part)
            for coordinate, value in cells.items():
                xml = patch_numeric(xml, coordinate, value)
            replacements[part] = xml
        with ZipFile(output, "w") as dst:
            dst.comment = src.comment
            for info in src.infolist():
                dst.writestr(info, replacements.get(info.filename, src.read(info.filename)))


def edit(source, output):
    with ZipFile(source) as archive:
        xml = archive.read(sheet_paths(archive)["Assumptions"])
        sheet = ET.fromstring(xml)
        previous = float(sheet.find(".//{%s}c[@r='B1']/{%s}v" % (MAIN, MAIN)).text)
    save_patches(source, output, {"Assumptions": {"B1": 0.12}})
    return {"changed_cell": "Assumptions!B1", "previous_value": previous,
            "new_value": 0.12, "recalculation": {"performed": False,
            "needed": True, "reason": "Existing formula caches were preserved; dependent formulas need recalculation."}}


def aggregate(source, output):
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        labels = [row[0] for row in book["Summary"].iter_rows(min_row=2, max_row=5, max_col=1, values_only=True)]
        totals = dict.fromkeys(labels, 0)
        selected = ["Period01", "Period03"]
        for title in selected:
            sheet = book[title]
            header = next(sheet.iter_rows(max_row=1, values_only=True))
            region_col, revenue_col = header.index("Region") + 1, header.index("Revenue") + 1
            first, last = min(region_col, revenue_col), max(region_col, revenue_col)
            for row in sheet.iter_rows(min_row=2, min_col=first, max_col=last, values_only=True):
                region, revenue = row[region_col-first], row[revenue_col-first]
                if region in totals and isinstance(revenue, (int, float)):
                    totals[region] += revenue
    finally:
        book.close()
    save_patches(source, output, {"Summary": {"B%d" % row: totals[label] for row, label in enumerate(labels, 2)}})
    return {"totals": totals, "selected_sheets": selected}


def wide(source, output):
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = book["Wide"]
        rows = sheet.iter_rows(min_row=1, max_row=41, values_only=True)
        header = next(rows)
        region_col, amount_col = header.index("Region"), header.index("Amount")
        total = sum(row[amount_col] for row in rows if row[region_col] == "APAC" and isinstance(row[amount_col], (int, float)))
    finally:
        book.close()
    save_patches(source, output, {"Summary": {"B2": total}})
    return {"total": total, "source_coordinates": {
        "region": "Wide!%s2:%s41" % (get_column_letter(region_col+1), get_column_letter(region_col+1)),
        "amount": "Wide!%s2:%s41" % (get_column_letter(amount_col+1), get_column_letter(amount_col+1))}}


def create(source, output):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Budget"
    for row in [["Month", "Revenue", "Cost", "Profit"], ["Jan", 100, 60, "=B2-C2"],
                ["Feb", 120, 70, "=B3-C3"], ["Mar", 90, 50, "=B4-C4"],
                ["Total", "=SUM(B2:B4)", "=SUM(C2:C4)", "=SUM(D2:D4)"]]:
        sheet.append(row)
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=11, bold=(cell.row == 1))
    for row in sheet.iter_rows(min_row=2, max_row=5, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = "0.00"
            if cell.data_type != "f":
                cell.comment = Comment("Source: values supplied in the task instruction.", "Source")
    sheet.freeze_panes = "A2"
    for column in "ABCD":
        sheet.column_dimensions[column].width = 16
    chart = BarChart()
    chart.type = "col"
    chart.title = "Monthly Revenue and Cost"
    chart.y_axis.title = "Amount"
    chart.x_axis.title = "Month"
    chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
    sheet.add_chart(chart, "F2")
    book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
    book.save(output)
    return {"sheet": "Budget", "range": "A1:D5", "formulas": 6,
            "chart": "Column chart of monthly Revenue and Cost, excluding totals",
            "formatting": "Bold Arial header, two decimals in B2:D5, frozen header row, readable widths",
            "recalculation": "Excel calculates on opening; formula caches omitted as requested"}


def main():
    task, source, output, result = sys.argv[1:]
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    Path(result).parent.mkdir(parents=True, exist_ok=True)
    summary = {"edit": edit, "aggregate": aggregate, "wide": wide, "create": create}[task](source, output)
    encoded = json.dumps(summary, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    if len(encoded) > 6000:
        raise ValueError("Result exceeds the requested 6000-byte limit")
    Path(result).write_bytes(encoded)


if __name__ == "__main__":
    main()
