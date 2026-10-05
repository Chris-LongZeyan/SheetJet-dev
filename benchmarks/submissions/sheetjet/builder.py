"""Frozen SheetJet task builders. New-workbook creation uses openpyxl fallback."""
import json
import os
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parent
TEMP = ROOT / "temp"
TEMP.mkdir(exist_ok=True)
os.environ["TEMP"] = str(TEMP)
os.environ["TMP"] = str(TEMP)


def patch(book, output, cells):
    ops = []
    for sheet, coord, value in cells:
        old = book.read_cells(sheet, [coord])[0]
        op = {"operation": "SET_VALUE", "sheet": sheet, "cell": coord,
              "value": value, "expected": old["value"]}
        if "formula" in old:
            op["expected_formula"] = old["formula"]
        ops.append(op)
    return book.patch_cells(ops, output)


def run_existing(task, source, output):
    from sheetjet import Workbook
    with Workbook(source, cache_dir=Path(output).parent / "sheetjet-cache") as book:
        if task == "edit":
            old = book.read_cells("Assumptions", ["B1"])[0]
            assert old["value"] == 0.08 and "formula" not in old
            report = book.patch_cells([{
                "operation": "SET_VALUE", "sheet": "Assumptions", "cell": "B1",
                "expected": 0.08, "value": 0.12,
            }], output)
            return {"changed_cell": "Assumptions!B1", "previous_value": 0.08,
                    "new_value": 0.12, "recalculated": False,
                    "requires_recalculation": report["requires_recalculation"],
                    "recalculation_status": "Requested on opening; existing formulas were not recalculated."}
        if task == "aggregate":
            sheets = ["Period01", "Period03"]
            book.query_engine.load_sheets(
                "SelectedRevenue", {s: "A1:E25001" for s in sheets},
                {"Region": "VARCHAR", "Revenue": "DECIMAL(18,2)"})
            rows = book.query_engine.query(
                'SELECT "Region", SUM("Revenue") FROM SelectedRevenue GROUP BY "Region"')['rows']
            totals = {region: float(value) for region, value in rows}
            labels = [r[0] for r in book.read_range("Summary", "A2:A5")["rows"]]
            assert len(set(labels)) == 4 and set(labels) == set(totals)
            patch(book, output, [("Summary", f"B{i}", totals[label])
                                 for i, label in enumerate(labels, 2)])
            return {"selected_sheets": sheets, "totals": {label: totals[label] for label in labels},
                    "destination": "Summary!B2:B5", "values": "static snapshots",
                    "recalculated": False, "requires_recalculation": True}
        if task == "wide":
            book.query_engine.load(
                "WideAmounts", sheet="Wide", ref="A1:CRD41", columns=["Region", "Amount"],
                schema={"Region": "VARCHAR", "Amount": "DECIMAL(18,2)"})
            value = book.query_engine.query(
                'SELECT SUM("Amount") FROM WideAmounts WHERE "Region" = ?', ["APAC"])["rows"][0][0]
            total = float(value)
            patch(book, output, [("Summary", "B2", total)])
            return {"total": total, "source_coordinates": {
                "Region": "Wide!A2:A41", "Amount": "Wide!CRD2:CRD41"}}
    raise ValueError(task)


def create(output):
    # SheetJet cannot create workbook structures or charts; no existing file is resaved here.
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font
    from openpyxl.workbook.properties import CalcProperties
    book = Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.append(["Month", "Revenue", "Cost", "Profit"])
    for i, row in enumerate([("Jan", 100, 60), ("Feb", 120, 70), ("Mar", 90, 50)], 2):
        sheet.append([*row, f"=B{i}-C{i}"])
    sheet.append(["Total", "=SUM(B2:B4)", "=SUM(C2:C4)", "=SUM(D2:D4)"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2, max_row=5, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = "0.00"
    sheet.freeze_panes = "A2"
    for column in "ABCD":
        sheet.column_dimensions[column].width = 16
    chart = BarChart()
    chart.type = "col"
    chart.title = "Monthly Revenue and Cost"
    chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
    sheet.add_chart(chart, "F2")
    book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    book.save(output)
    return {"sheet": "Budget", "range": "A1:D5", "months": ["Jan", "Feb", "Mar"],
            "formulas": ["D2:D4: monthly profit", "B5:D5: column totals"],
            "chart": "Column chart of monthly Revenue and Cost; totals excluded",
            "formatting": "Bold header; B2:D5 two decimals; frozen header; 16-character column widths",
            "engine": "openpyxl fallback: SheetJet does not create workbooks or charts",
            "formula_caches": "Not calculated; Excel recalculation requested on opening"}


def main():
    task, source, output, result = sys.argv[1:]
    output, result = Path(output), Path(result)
    if not output.is_absolute() or not result.is_absolute():
        raise ValueError("Output and result paths must be absolute")
    output.parent.mkdir(parents=True, exist_ok=True)
    result.parent.mkdir(parents=True, exist_ok=True)
    data = create(output) if task == "create" else run_existing(task, source, output)
    encoded = json.dumps(data, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    assert len(encoded) <= 6000
    result.write_bytes(encoded)


if __name__ == "__main__":
    main()
