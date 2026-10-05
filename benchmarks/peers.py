"""Same fixtures, verified answers, isolated processes, actual peer libraries.

Cold and cached workloads are separate rows. No peer is penalized for returning
bounded answers; context payload sizes are identical when answers are identical.
"""

from __future__ import annotations

import argparse
import gc
import hashlib
import importlib.metadata
import importlib.util
import json
import platform
import random
import statistics
import subprocess
import sys
import tempfile
import threading
import time
from collections import defaultdict
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

METHODS = [
    "sheetjet",
    "openpyxl",
    "pandas_openpyxl",
    "pandas_calamine",
    "polars_calamine",
    "calamine",
    "sheetjet_cached",
    "polars_parquet",
    "sheetjet_session",
    "polars_session",
    "sheetjet_materialized",
    "polars_materialized",
    "sheetjet_patch",
    "openpyxl_patch",
]


def fixture(path, rows, sheets, shared_strings=False):
    """Each selected sheet has the same deterministic data; unrelated sheets are included."""
    if not 1 <= rows <= 1048575 or not 1 <= sheets <= 100:
        raise ValueError("Invalid fixture dimensions")
    names = [f"Period{i + 1:02}" for i in range(sheets)]
    if shared_strings:
        import xlsxwriter

        with xlsxwriter.Workbook(path) as b:
            for name in names:
                s = b.add_worksheet(name)
                s.write_row(0, 0, ["ID", "Customer", "Region", "Revenue", "Cost"])
                for i in range(1, rows + 1):
                    s.write_row(
                        i,
                        0,
                        [
                            i,
                            f"Customer{i % 1000}",
                            ["APAC", "EMEA", "Americas", "Other"][i % 4],
                            i % 1000 + 1,
                            i % 500,
                        ],
                    )
            a = b.add_worksheet("Assumptions")
            a.write_row(0, 0, ["Growth", 0.08])
            a = b.add_worksheet("HiddenNotes")
            a.hide()
            a.write_row(0, 0, ["Label", "Value"])
            a.write_row(1, 0, ["Excluded from selected ranges", 999999999])
        return names
    from openpyxl import Workbook

    from .fixtures import transactions

    with tempfile.TemporaryDirectory(dir=path.parent) as tmp:
        prototype = transactions(Path(tmp) / "prototype.xlsx", rows)
        b = Workbook()
        b.remove(b.active)
        for name in names:
            s = b.create_sheet(name)
            s.append(["ID", "Customer", "Region", "Revenue", "Cost"])
        b.create_sheet("Assumptions").append(["Growth", 0.08])
        hidden = b.create_sheet("HiddenNotes")
        hidden.sheet_state = "hidden"
        hidden.append(["Label", "Value"])
        hidden.append(["Excluded from selected ranges", 999999999])
        template = Path(tmp) / "template.xlsx"
        b.save(template)
        with ZipFile(prototype) as z:
            sheet = z.read("xl/worksheets/sheet1.xml")
        # The union API uses explicit ranges, so no table relationship is needed.
        sheet = sheet[: sheet.index(b"<tableParts")] + b"</worksheet>"
        with (
            ZipFile(template) as src,
            ZipFile(path, "w", compression=ZIP_DEFLATED, compresslevel=1) as dst,
        ):
            targets = {f"xl/worksheets/sheet{i + 1}.xml" for i in range(sheets)}
            for info in src.infolist():
                if info.filename in targets:
                    dst.writestr(info.filename, sheet)
                else:
                    dst.writestr(info, src.read(info.filename))
    return names


def _sql():
    return 'SELECT "Region",SUM("Revenue") AS revenue FROM Sales GROUP BY 1 ORDER BY 1'


def worker(method, path, rows, sheets):
    import psutil

    process = psutil.Process()
    peak = [process.memory_info().rss]
    stop = threading.Event()

    def sample():
        while not stop.wait(0.01):
            peak[0] = max(peak[0], process.memory_info().rss)

    thread = threading.Thread(target=sample, daemon=True)
    thread.start()
    names = [f"Period{i + 1:02}" for i in range(sheets)]
    ranges = {n: f"A1:E{rows + 1}" for n in names}
    schema = {"Region": "VARCHAR", "Revenue": "BIGINT"}
    metrics = {}
    setup = 0.0
    latencies, repeated_answers = [], []
    try:
        with tempfile.TemporaryDirectory(prefix="sheetjet-peer-") as tmp:
            tmp = Path(tmp)
            start = time.perf_counter()
            if method == "sheetjet_cached":
                from sheetjet import Workbook

                with Workbook(path, tmp / "cache") as w:
                    w.query_engine.load_sheets("Sales", ranges, schema)
                setup = time.perf_counter() - start
                gc.collect()
                start = time.perf_counter()
            elif method in {"polars_parquet", "polars_session", "polars_materialized"}:
                import polars as pl

                frames = pl.read_excel(
                    path,
                    sheet_name=names,
                    columns=["Region", "Revenue"],
                    schema_overrides={"Region": pl.String, "Revenue": pl.Int64},
                )
                for i, frame in enumerate(frames.values()):
                    frame.write_parquet(tmp / f"{i}.parquet", compression="zstd")
                del frames, frame
                setup = time.perf_counter() - start
                gc.collect()
                start = time.perf_counter()
            if method in {
                "sheetjet",
                "sheetjet_cached",
                "sheetjet_session",
                "sheetjet_materialized",
            }:
                from sheetjet import Workbook

                with Workbook(path, tmp / "cache") as w:
                    w.query_engine.load_sheets("Sales", ranges, schema)
                    if method == "sheetjet_materialized":
                        w.query_engine.materialize("Sales")
                    if method in {"sheetjet_session", "sheetjet_materialized"}:
                        w.query_engine.query(
                            _sql()
                        )  # Establish a warm query session for both peers.
                        setup = time.perf_counter() - start
                        gc.collect()
                        start = time.perf_counter()
                    for _ in range(
                        20 if method in {"sheetjet_session", "sheetjet_materialized"} else 1
                    ):
                        query_start = time.perf_counter()
                        answer = w.query_engine.query(_sql())["rows"]
                        latencies.append(time.perf_counter() - query_start)
                        repeated_answers.append(answer)
                    metrics = w.metrics.snapshot()
            elif method.startswith("pandas_"):
                import pandas as pd

                engine = method.removeprefix("pandas_")
                frames = pd.read_excel(
                    path,
                    sheet_name=names,
                    usecols=["Region", "Revenue"],
                    engine=engine,
                    dtype={"Region": str, "Revenue": "int64"},
                )
                totals = defaultdict(int)
                for frame in frames.values():
                    for region, value in frame.groupby("Region")["Revenue"].sum().items():
                        totals[region] += int(value)
                answer = [[k, v] for k, v in sorted(totals.items())]
            elif method in {
                "polars_calamine",
                "polars_parquet",
                "polars_session",
                "polars_materialized",
            }:
                import polars as pl

                if method in {"polars_parquet", "polars_session", "polars_materialized"}:
                    source = pl.scan_parquet(str(tmp / "*.parquet"))
                    if method == "polars_materialized":
                        source = source.collect().lazy()
                    lazy = source.group_by("Region").agg(pl.col("Revenue").sum()).sort("Region")
                    if method in {"polars_session", "polars_materialized"}:
                        lazy.collect()
                        setup += time.perf_counter() - start
                        gc.collect()
                        start = time.perf_counter()
                    for _ in range(
                        20 if method in {"polars_session", "polars_materialized"} else 1
                    ):
                        query_start = time.perf_counter()
                        data = lazy.collect()
                        latencies.append(time.perf_counter() - query_start)
                        repeated_answers.append([list(row) for row in data.rows()])
                else:
                    frames = pl.read_excel(
                        path,
                        sheet_name=names,
                        columns=["Region", "Revenue"],
                        schema_overrides={"Region": pl.String, "Revenue": pl.Int64},
                    )
                    data = (
                        pl.concat(
                            [
                                f.group_by("Region").agg(pl.col("Revenue").sum())
                                for f in frames.values()
                            ]
                        )
                        .group_by("Region")
                        .agg(pl.col("Revenue").sum())
                        .sort("Region")
                    )
                answer = [list(row) for row in data.rows()]
            elif method == "calamine":
                from python_calamine import CalamineWorkbook

                totals = defaultdict(int)
                with CalamineWorkbook.from_path(str(path)) as b:
                    for name in names:
                        iterator = iter(b.get_sheet_by_name(name).iter_rows())
                        next(iterator)
                        for row in iterator:
                            totals[row[2]] += int(row[3])
                answer = [[k, v] for k, v in sorted(totals.items())]
            elif method == "openpyxl":
                from openpyxl import load_workbook

                totals = defaultdict(int)
                b = load_workbook(path, read_only=True, data_only=True)
                try:
                    for name in names:
                        for region, revenue in b[name].iter_rows(
                            min_row=2, max_row=rows + 1, min_col=3, max_col=4, values_only=True
                        ):
                            totals[region] += revenue
                finally:
                    b.close()
                answer = [[k, v] for k, v in sorted(totals.items())]
            elif method == "sheetjet_patch":
                from sheetjet import Workbook

                with Workbook(path, tmp / "cache") as w:
                    w.patch_cells(
                        [
                            {
                                "operation": "SET_VALUE",
                                "sheet": "Assumptions",
                                "cell": "B1",
                                "expected": 0.08,
                                "value": 0.12,
                            }
                        ],
                        tmp / "edited.xlsx",
                    )
                    metrics = w.metrics.snapshot()
                answer = 0.12
            elif method == "openpyxl_patch":
                from openpyxl import load_workbook

                b = load_workbook(path)
                b["Assumptions"]["B1"] = 0.12
                b.save(tmp / "edited.xlsx")
                b.close()
                answer = 0.12
            else:
                raise ValueError(method)
            elapsed = time.perf_counter() - start
            if method.endswith("_patch"):
                # Identical independent saved-value check, excluded from each timed region.
                from sheetjet import Workbook

                with Workbook(tmp / "edited.xlsx", tmp / "verify") as w:
                    assert w.read_range("Assumptions", "B1")["rows"] == [[0.12]]
            else:
                totals = defaultdict(int)
                for i in range(1, rows + 1):
                    totals[["APAC", "EMEA", "Americas", "Other"][i % 4]] += (i % 1000 + 1) * sheets
                assert answer == [[k, v] for k, v in sorted(totals.items())], answer
                assert all(result == answer for result in repeated_answers)
    finally:
        stop.set()
        thread.join()
    return {
        "method": method,
        "seconds": elapsed,
        "setup_seconds": setup,
        "peak_rss_mb": peak[0] / 1024**2,
        "answer": answer,
        "verified": True,
        "metrics": metrics,
        "query_repetitions": len(latencies),
        "query_latencies_seconds": latencies,
    }


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--rows", type=int, default=25000, help="Data rows per sheet")
    p.add_argument("--sheets", type=int, default=4)
    p.add_argument("--repeats", type=int, default=3)
    p.add_argument("--output", type=Path, default=Path("benchmark-output/peers/results.json"))
    p.add_argument("--shared-strings", action="store_true")
    p.add_argument("--workbook", type=Path)
    p.add_argument("--methods", nargs="+", choices=METHODS, default=METHODS)
    p.add_argument("--worker", choices=METHODS)
    args = p.parse_args()
    if args.worker:
        print(
            json.dumps(
                worker(args.worker, args.workbook, args.rows, args.sheets), separators=(",", ":")
            )
        )
        return
    args.output.parent.mkdir(parents=True, exist_ok=True)
    path = args.workbook or args.output.parent / "multisheet.xlsx"
    if not args.workbook:
        fixture(path, args.rows, args.sheets, args.shared_strings)
    jobs = [(method, i) for i in range(args.repeats) for method in args.methods]
    random.Random(20261005).shuffle(jobs)
    trials = []
    for method, i in jobs:
        result = subprocess.run(
            [
                sys.executable,
                "-m",
                "benchmarks.peers",
                "--worker",
                method,
                "--rows",
                str(args.rows),
                "--sheets",
                str(args.sheets),
                "--workbook",
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=600,
        )
        if result.returncode:
            raise RuntimeError(f"{method}: {result.stderr[-4000:]}")
        trial = json.loads(result.stdout)
        trials.append(trial)
        print(
            f"{method} {i + 1}/{args.repeats}: {trial['seconds']:.3f}s, {trial['peak_rss_mb']:.1f} MiB",
            flush=True,
        )
    versions = {
        m: importlib.metadata.version(m)
        for m in ["duckdb", "lxml", "openpyxl", "polars", "fastexcel", "pandas", "python-calamine"]
    }
    with path.open("rb") as stream:
        digest = hashlib.file_digest(stream, "sha256").hexdigest()
    package_root = Path(importlib.util.find_spec("sheetjet").submodule_search_locations[0])
    source_hash = hashlib.sha256()
    for source in sorted(package_root.glob("*.py")):
        source_hash.update(source.name.encode() + b"\0" + source.read_bytes())
    result = {
        "rows_per_sheet": args.rows,
        "selected_sheets": args.sheets,
        "total_rows": args.rows * args.sheets,
        "shared_strings": args.shared_strings,
        "repeats": args.repeats,
        "file_bytes": path.stat().st_size,
        "file_sha256": digest,
        "sheetjet_source_sha256": source_hash.hexdigest(),
        "benchmark_source_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "session_methodology": "Methods ending in _session or _materialized perform 20 identical aggregations after one untimed warm-up; seconds is the whole 20-query phase, not one query. Individual latencies are retained. Materialized routes copy rows into a session DuckDB table or Polars DataFrame before timing. Preparation, including that copy, is reported separately. Every repeated answer is checked; no answer cache is used.",
        "environment": {
            "python": platform.python_version(),
            "platform": platform.platform(),
            "cpu": platform.processor(),
            "versions": versions,
        },
        "methodology": "Fresh process per trial, fixed randomized order; OS cache not flushed. Cold time includes method imports, load, projection, aggregation, close; excludes fixture creation, temporary-directory cleanup and independent validation. Cached rows exclude preparation/imports and measure a new SheetJet workbook session or Polars Parquet query; setup time is reported separately. RSS sampled every 10ms includes setup/validation and can miss spikes. All methods use the same selected sheets and aggregate inputs; returned answers match independent arithmetic. Cache comparison includes a peer Parquet workflow. Patch tasks edit the small Assumptions sheet, not a large sheet. No native Excel verification. No claim of universal superiority.",
        "medians": [
            {
                "method": m,
                **{
                    k: statistics.median(t[k] for t in trials if t["method"] == m)
                    for k in ["seconds", "setup_seconds", "peak_rss_mb"]
                },
            }
            for m in args.methods
        ],
        "trials": trials,
    }
    args.output.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(f"Saved {args.output}")


if __name__ == "__main__":
    main()
